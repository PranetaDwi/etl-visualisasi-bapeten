
import pandas as pd
import csv
import io
from openpyxl import load_workbook
import sys
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent

if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from services.google_services_own_account import (
    connect_google,
    read_excel,
    write_sheet,
    upload_parquet,
)

def refresh_data_srp():

    # 1. CONNECT SERVICES
    drive, sheets = connect_google()

    # 2. GET DATA EXCEL
    FILE_EXCEL_ID = "1rxJcxH482wk-Z7H3oW76LqzrfVS_H1hf"

    file_buffer_excel = read_excel(
        FILE_EXCEL_ID,
        drive
    )

    # 3. PARSING DATA
    ws = load_workbook(
        file_buffer_excel,
        read_only=True,
        data_only=True
    ).active

    try:
        text = "\n".join(
            str(r[0])
            for r in ws.iter_rows(values_only=True)
            if r and r[0] is not None
        )
    finally:
        ws.parent.close()

    rows = list(csv.reader(io.StringIO(text)))

    if not rows:
        raise ValueError("File Excel tidak memiliki data.")

    header, data = rows[0], rows[1:]

    good = [
        r for r in data
        if len(r) == len(header)
    ]

    bad = [
        r for r in data
        if len(r) != len(header)
    ]

    df = pd.DataFrame(
        good,
        columns=header
    ).replace("", pd.NA)

    # 4. CLEANING DATA
    df_valid = df.drop_duplicates().copy()

    df_valid = df_valid[
        df_valid["no_seri"].notna()
        & df_valid["no_seri"].astype(str).str.strip().ne("")
    ]

    df_valid = df_valid[
        df_valid["propinsi"].notna()
        & df_valid["propinsi"].astype(str).str.strip().ne("")
    ]

    # 5. REKAPITULASI PER PROVINSI
    df_total = (
        df_valid
        .groupby(
            ["propinsi", "jenis_sumber"],
            dropna=False
        )
        .size()
        .unstack(fill_value=0)
        .reindex(
            columns=["PRP", "ZRA"],
            fill_value=0
        )
    )

    df_total["Total"] = (
        df_total["PRP"] + df_total["ZRA"]
    )

    df_total = (
        df_total
        .rename(columns={
            "PRP": "Jumlah PRP",
            "ZRA": "Jumlah ZRA",
            "Total": "Total SRP"
        })
        .reset_index()
        .rename(columns={
            "propinsi": "Provinsi"
        })
    )

    # 6. WRITE TO GOOGLE SHEETS
    FILE_REKAPITULASI_ID = (
        "1jZt22z352D151rMTqZITiPbenX45VIC7iTWe79VyjsA"
    )

    write_sheet(
        df_total,
        FILE_REKAPITULASI_ID,
        range_name="Sheet1!B1",
        sheets=sheets
    )

    # 7. UPLOAD DATA BERSIH KE GOOGLE DRIVE
    PARQUET_FILE_ID = upload_parquet(
        df=df_valid,
        drive=drive,
        filename="data_srp_clean.parquet",
        folder_id="1Xzl-d3y11wCdr21yjp1kdhrUkY96QJ2X"
    )

    print(f"Jumlah data valid: {len(df_valid)}")
    print(f"Jumlah baris tidak valid: {len(bad)}")
    print(f"Parquet tersimpan dengan ID: {PARQUET_FILE_ID}")

    # 8. RETURN DATA BERSIH
    return df_valid

refresh_data_srp()